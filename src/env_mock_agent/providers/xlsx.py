from __future__ import annotations

from datetime import datetime, timedelta

from openpyxl import Workbook

from env_mock_agent.providers.base import ArtifactProvider, ProviderRequest
from env_mock_agent.providers.helpers import artifact_result, normalize_zip_timestamps
from env_mock_agent.schemas import ArtifactResult


class XlsxProvider(ArtifactProvider):
    name = "xlsx"
    asset_types = ("xlsx",)

    def generate(self, request: ProviderRequest) -> ArtifactResult:
        target = request.destination()
        contract = request.plan.content_contract
        workbook = Workbook()
        active = workbook.active
        if active is not None:
            workbook.remove(active)
        sheets = contract.get("sheets")
        normalized_sheets = sheets if isinstance(sheets, list) else []
        if not normalized_sheets:
            normalized_sheets = [{"name": "Data", "rows": contract.get("rows", [])}]
        for index, raw_sheet in enumerate(normalized_sheets, start=1):
            if not isinstance(raw_sheet, dict):
                continue
            name = str(raw_sheet.get("name") or f"Sheet{index}")[:31]
            worksheet = workbook.create_sheet(name)
            rows = raw_sheet.get("rows")
            if isinstance(rows, list):
                for raw_row in rows:
                    if isinstance(raw_row, list):
                        worksheet.append(raw_row)
                    elif isinstance(raw_row, dict):
                        worksheet.append(list(raw_row.values()))
            freeze_panes = raw_sheet.get("freeze_panes")
            if isinstance(freeze_panes, str):
                worksheet.freeze_panes = freeze_panes
            widths = raw_sheet.get("column_widths")
            if isinstance(widths, dict):
                for column, width in widths.items():
                    if isinstance(width, int | float):
                        worksheet.column_dimensions[str(column)].width = float(width)
        if not workbook.sheetnames:
            workbook.create_sheet("Data")
        workbook.properties.creator = str(contract.get("author") or "")
        workbook.properties.lastModifiedBy = ""
        workbook.properties.title = str(contract.get("title") or "")
        workbook.properties.subject = str(contract.get("subject") or "")
        document_date = datetime(2020, 1, 1) + timedelta(days=request.plan.seed % 366)
        workbook.properties.created = document_date
        workbook.properties.modified = document_date
        workbook.save(target)
        normalize_zip_timestamps(target, request.plan.seed)
        return artifact_result(
            request.plan.artifact_id,
            target,
            self.name,
            metadata={"sheets": workbook.sheetnames},
        )
