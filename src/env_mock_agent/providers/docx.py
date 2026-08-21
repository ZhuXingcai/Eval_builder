from __future__ import annotations

from datetime import datetime, timedelta

from docx import Document

from env_mock_agent.providers.base import ArtifactProvider, ProviderRequest
from env_mock_agent.providers.helpers import (
    artifact_result,
    as_string_list,
    normalize_zip_timestamps,
)
from env_mock_agent.schemas import ArtifactResult


class DocxProvider(ArtifactProvider):
    name = "docx"
    asset_types = ("docx",)

    def generate(self, request: ProviderRequest) -> ArtifactResult:
        target = request.destination()
        contract = request.plan.content_contract
        document = Document()
        title = str(contract.get("title") or "")
        if title:
            document.add_heading(title, level=0)
        for paragraph in as_string_list(contract.get("paragraphs")):
            document.add_paragraph(paragraph)
        sections = contract.get("sections")
        if isinstance(sections, list):
            for raw_section in sections:
                if not isinstance(raw_section, dict):
                    continue
                heading = str(raw_section.get("heading") or "")
                if heading:
                    document.add_heading(heading, level=1)
                content = raw_section.get("paragraphs", raw_section.get("content"))
                for paragraph in as_string_list(content):
                    document.add_paragraph(paragraph)
        tables = contract.get("tables")
        if isinstance(tables, list):
            for raw_table in tables:
                if not isinstance(raw_table, dict):
                    continue
                headers = as_string_list(raw_table.get("headers"))
                rows = raw_table.get("rows")
                normalized_rows = rows if isinstance(rows, list) else []
                width = max(
                    len(headers),
                    max(
                        (len(row) for row in normalized_rows if isinstance(row, list)),
                        default=0,
                    ),
                )
                if width == 0:
                    continue
                table = document.add_table(rows=1 if headers else 0, cols=width)
                if headers:
                    for index, value in enumerate(headers):
                        table.rows[0].cells[index].text = value
                for raw_row in normalized_rows:
                    if not isinstance(raw_row, list):
                        continue
                    cells = table.add_row().cells
                    for index, value in enumerate(raw_row[:width]):
                        cells[index].text = str(value)
        properties = document.core_properties
        properties.author = str(contract.get("author") or "")
        properties.last_modified_by = ""
        properties.title = title
        properties.subject = str(contract.get("subject") or "")
        properties.keywords = ""
        properties.comments = ""
        document_date = datetime(2020, 1, 1) + timedelta(days=request.plan.seed % 366)
        properties.created = document_date
        properties.modified = document_date
        document.save(str(target))
        normalize_zip_timestamps(target, request.plan.seed)
        return artifact_result(
            request.plan.artifact_id,
            target,
            self.name,
            metadata={"paragraphs": len(document.paragraphs), "tables": len(document.tables)},
        )
