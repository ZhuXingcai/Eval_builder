from __future__ import annotations

from docx import Document

from env_mock_agent.schemas import FindingCategory, Severity, ValidationFinding
from env_mock_agent.validators.base import (
    ArtifactValidator,
    ValidationRequest,
    finding,
)


class DocxValidator(ArtifactValidator):
    name = "docx"

    def validate(self, request: ValidationRequest) -> list[ValidationFinding]:
        path = request.artifact_path()
        try:
            document = Document(str(path))
        except Exception as exc:
            return [
                finding(
                    self.name,
                    request,
                    Severity.P0,
                    FindingCategory.STRUCTURE,
                    f"DOCX cannot be parsed: {exc}",
                    repair_action="regenerate with the DOCX provider",
                )
            ]
        contract = request.plan.content_contract
        findings: list[ValidationFinding] = []
        minimum_paragraphs = contract.get("min_paragraphs", 0)
        if isinstance(minimum_paragraphs, int) and len(document.paragraphs) < minimum_paragraphs:
            findings.append(
                finding(
                    self.name,
                    request,
                    Severity.P1,
                    FindingCategory.COVERAGE,
                    f"DOCX has {len(document.paragraphs)} paragraphs; expected {minimum_paragraphs}",
                )
            )
        minimum_tables = contract.get("min_tables", 0)
        if isinstance(minimum_tables, int) and len(document.tables) < minimum_tables:
            findings.append(
                finding(
                    self.name,
                    request,
                    Severity.P1,
                    FindingCategory.COVERAGE,
                    f"DOCX has {len(document.tables)} tables; expected {minimum_tables}",
                )
            )
        if not any(paragraph.text.strip() for paragraph in document.paragraphs) and not document.tables:
            findings.append(
                finding(
                    self.name,
                    request,
                    Severity.P0,
                    FindingCategory.COVERAGE,
                    "DOCX contains no readable paragraphs or tables",
                )
            )
        return findings
