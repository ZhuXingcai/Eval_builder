from __future__ import annotations

import pymupdf
from docx import Document
from openpyxl import load_workbook

from env_mock_agent.schemas import FindingCategory, Severity, ValidationFinding
from env_mock_agent.validators.base import (
    ArtifactValidator,
    ValidationRequest,
    finding,
)

AUTOMATION_MARKERS = ("python-docx", "openpyxl", "reportlab")


class MetadataValidator(ArtifactValidator):
    name = "metadata"

    def validate(self, request: ValidationRequest) -> list[ValidationFinding]:
        path = request.artifact_path()
        suffix = path.suffix.lower()
        values: list[str] = []
        if suffix == ".docx":
            properties = Document(str(path)).core_properties
            values.extend(
                str(value or "")
                for value in (
                    properties.author,
                    properties.last_modified_by,
                    properties.comments,
                )
            )
        elif suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True)
            try:
                values.extend(
                    str(value or "")
                    for value in (
                        workbook.properties.creator,
                        workbook.properties.lastModifiedBy,
                        workbook.properties.description,
                    )
                )
            finally:
                workbook.close()
        elif suffix == ".pdf":
            with pymupdf.open(path) as document:  # type: ignore[no-untyped-call]
                values.extend(str(value or "") for value in document.metadata.values())
        combined = "\n".join(values).casefold()
        findings: list[ValidationFinding] = []
        for marker in AUTOMATION_MARKERS:
            if marker in combined:
                findings.append(
                    finding(
                        self.name,
                        request,
                        Severity.P1,
                        FindingCategory.METADATA,
                        f"automation marker remains in document metadata: {marker}",
                        repair_action="rewrite document core metadata before release",
                    )
                )
        return findings
