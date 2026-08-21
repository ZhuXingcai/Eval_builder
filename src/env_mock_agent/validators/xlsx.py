from __future__ import annotations

from openpyxl import load_workbook

from env_mock_agent.schemas import FindingCategory, Severity, ValidationFinding
from env_mock_agent.validators.base import (
    ArtifactValidator,
    ValidationRequest,
    finding,
)


class XlsxValidator(ArtifactValidator):
    name = "xlsx"

    def validate(self, request: ValidationRequest) -> list[ValidationFinding]:
        path = request.artifact_path()
        try:
            workbook = load_workbook(path, read_only=True, data_only=False)
        except Exception as exc:
            return [
                finding(
                    self.name,
                    request,
                    Severity.P0,
                    FindingCategory.STRUCTURE,
                    f"XLSX cannot be parsed: {exc}",
                    repair_action="regenerate with the XLSX provider",
                )
            ]
        findings: list[ValidationFinding] = []
        try:
            required_sheets = request.plan.content_contract.get("required_sheets")
            if isinstance(required_sheets, list):
                for sheet_name in required_sheets:
                    if str(sheet_name) not in workbook.sheetnames:
                        findings.append(
                            finding(
                                self.name,
                                request,
                                Severity.P1,
                                FindingCategory.COVERAGE,
                                f"required worksheet is missing: {sheet_name}",
                            )
                        )
            nonempty_cells = 0
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows(values_only=True):
                    nonempty_cells += sum(value is not None for value in row)
            if nonempty_cells == 0:
                findings.append(
                    finding(
                        self.name,
                        request,
                        Severity.P0,
                        FindingCategory.COVERAGE,
                        "XLSX contains no populated cells",
                    )
                )
        finally:
            workbook.close()
        return findings
