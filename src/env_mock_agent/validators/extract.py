from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import load_workbook

TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".html",
    ".htm",
    ".rtf",
    ".json",
    ".yaml",
    ".yml",
    ".xml",
    ".csv",
    ".tsv",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".go",
    ".rs",
    ".java",
    ".kt",
    ".sql",
    ".toml",
}


def extract_text(path: Path, *, max_characters: int = 2_000_000) -> str:
    if path.is_dir():
        chunks: list[str] = []
        for child in sorted(item for item in path.rglob("*") if item.is_file()):
            if child.suffix.lower() not in TEXT_SUFFIXES:
                continue
            chunks.append(child.read_text(encoding="utf-8", errors="replace"))
            if sum(len(chunk) for chunk in chunks) >= max_characters:
                break
        return "\n".join(chunks)[:max_characters]
    suffix = path.suffix.lower()
    if suffix == ".docx":
        document = Document(str(path))
        chunks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            chunks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(chunks)[:max_characters]
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=False)
        chunks = []
        try:
            for worksheet in workbook.worksheets:
                chunks.append(worksheet.title)
                for row in worksheet.iter_rows(values_only=True):
                    chunks.append(" | ".join("" if value is None else str(value) for value in row))
        finally:
            workbook.close()
        return "\n".join(chunks)[:max_characters]
    if suffix == ".pdf":
        with pymupdf.open(path) as document:  # type: ignore[no-untyped-call]
            return "\n".join(page.get_text() for page in document)[:max_characters]
    if suffix in TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8", errors="replace")[:max_characters]
    return ""
