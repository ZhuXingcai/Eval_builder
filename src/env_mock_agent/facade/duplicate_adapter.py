from __future__ import annotations

import hashlib
import io
import re
import stat
import unicodedata
import zipfile
from collections.abc import Mapping
from dataclasses import dataclass, field
from enum import StrEnum
from pathlib import Path, PurePosixPath
from typing import Protocol

import pymupdf
from docx import Document
from openpyxl import load_workbook

from env_mock_agent.facade.duplicate_v2 import (
    ATTACHMENT_DUPLICATE_FINGERPRINT_POLICY_VERSION,
    ATTACHMENT_DUPLICATE_NORMALIZATION_VERSION,
    AttachmentDuplicateFingerprintFacade,
    AttachmentDuplicateFingerprintFailureCodeV2,
    AttachmentDuplicateFingerprintOutcomeV2,
    AttachmentDuplicateFingerprintRequestV2,
    AttachmentDuplicateFingerprintResultV2,
    attachment_duplicate_fingerprint_request_ref,
    attachment_duplicate_fingerprint_result_carried_sha256,
    validate_attachment_duplicate_fingerprint_request_identity,
)
from env_mock_agent.providers.helpers import sha256_path
from env_mock_agent.runtimes.security import is_path_inside
from env_mock_agent.validators.extract import TEXT_SUFFIXES, extract_text

ATTACHMENT_DUPLICATE_EXTRACTOR_VERSION = "attachment-duplicate-extractor/r7-01-v1"

_TOKEN_PATTERN = re.compile(
    r"[\w]+(?:[._/-][\w]+)*",
    flags=re.UNICODE,
)
_TEXT_SUFFIXES = TEXT_SUFFIXES | {
    ".cfg",
    ".conf",
    ".env",
    ".ini",
    ".log",
    ".properties",
    ".sh",
}
_STRUCTURED_DOCUMENT_SUFFIXES = {".docx", ".xlsx", ".pdf"}
_UNSUPPORTED_CONTAINER_SUFFIXES = {
    ".7z",
    ".bz2",
    ".gz",
    ".rar",
    ".tar",
    ".tgz",
    ".xz",
}


class AttachmentTextExtractionLimits(Protocol):
    @property
    def max_file_bytes(self) -> int: ...

    @property
    def max_extracted_characters(self) -> int: ...

    @property
    def max_inventory_members(self) -> int: ...

    @property
    def max_nested_depth(self) -> int: ...

    @property
    def max_expanded_bytes(self) -> int: ...

    @property
    def max_compression_ratio_milli(self) -> int: ...


class AttachmentTextExtractionRequest(Protocol):
    @property
    def content_sha256(self) -> str: ...

    @property
    def size_bytes(self) -> int: ...

    @property
    def limits(self) -> AttachmentTextExtractionLimits: ...


class AttachmentTextExtractionFailureReason(StrEnum):
    CONTAINER_ENCRYPTED = "CONTAINER_ENCRYPTED"
    CONTAINER_TRUNCATED = "CONTAINER_TRUNCATED"


class AttachmentTextExtractionError(RuntimeError):
    def __init__(
        self,
        message: str,
        *,
        code: AttachmentDuplicateFingerprintFailureCodeV2,
        reason: AttachmentTextExtractionFailureReason | None = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.reason = reason


class AttachmentDuplicateFingerprintMaterialError(AttachmentTextExtractionError):
    pass


class _FingerprintBlocked(AttachmentTextExtractionError):
    def __init__(
        self,
        code: AttachmentDuplicateFingerprintFailureCodeV2,
        *,
        reason: AttachmentTextExtractionFailureReason | None = None,
    ) -> None:
        super().__init__(code.value, code=code, reason=reason)


class _FingerprintUnsupported(AttachmentTextExtractionError):
    def __init__(
        self,
        code: AttachmentDuplicateFingerprintFailureCodeV2,
    ) -> None:
        super().__init__(code.value, code=code)


class AttachmentDuplicateFingerprintMaterialResolver(Protocol):
    def resolve_output(
        self,
        request: AttachmentDuplicateFingerprintRequestV2,
    ) -> Path: ...


class MappingAttachmentDuplicateFingerprintMaterialResolver:
    def __init__(
        self,
        *,
        outputs: Mapping[str, Path],
    ) -> None:
        self._outputs = dict(outputs)

    def resolve_output(
        self,
        request: AttachmentDuplicateFingerprintRequestV2,
    ) -> Path:
        try:
            return self._outputs[request.output_ref.object_id].expanduser()
        except KeyError as exc:
            raise AttachmentDuplicateFingerprintMaterialError(
                "attachment output is not registered",
                code=(AttachmentDuplicateFingerprintFailureCodeV2.MATERIAL_NOT_FOUND),
            ) from exc


class StagingAttachmentDuplicateFingerprintMaterialResolver:
    def __init__(
        self,
        *,
        staging_root: Path,
        outputs: Mapping[str, Path],
    ) -> None:
        self._staging_root = staging_root.expanduser().resolve()
        self._outputs = dict(outputs)

    def resolve_output(
        self,
        request: AttachmentDuplicateFingerprintRequestV2,
    ) -> Path:
        try:
            target = self._outputs[request.output_ref.object_id].expanduser()
        except KeyError as exc:
            raise AttachmentDuplicateFingerprintMaterialError(
                "attachment output is not registered",
                code=(AttachmentDuplicateFingerprintFailureCodeV2.MATERIAL_NOT_FOUND),
            ) from exc
        resolved = target.resolve(strict=False)
        if not is_path_inside(self._staging_root, resolved):
            raise AttachmentDuplicateFingerprintMaterialError(
                "attachment output escapes staging",
                code=(AttachmentDuplicateFingerprintFailureCodeV2.OUTPUT_PATH_VIOLATION),
            )
        return target


@dataclass(slots=True)
class _ExtractionState:
    request: AttachmentTextExtractionRequest
    require_all_members_scannable: bool = False
    chunks: list[str] = field(default_factory=list)
    paths: list[str] = field(default_factory=list)
    path_set: set[str] = field(default_factory=set)
    inventory_members: int = 0
    expanded_bytes: int = 0
    extracted_characters: int = 0

    def add_member(
        self,
        *,
        size_bytes: int,
        logical_path: str | None = None,
    ) -> None:
        self.inventory_members += 1
        if self.inventory_members > self.request.limits.max_inventory_members:
            raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.SCAN_LIMIT_EXCEEDED)
        self.expanded_bytes += size_bytes
        if self.expanded_bytes > self.request.limits.max_expanded_bytes:
            raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.SCAN_LIMIT_EXCEEDED)
        if logical_path is not None:
            self.add_path(logical_path)

    def add_path(self, logical_path: str) -> None:
        if logical_path in self.path_set:
            raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE)
        self.path_set.add(logical_path)
        self.paths.append(logical_path)

    def add_text(self, text: str) -> None:
        observed = self.extracted_characters + len(text)
        if observed > self.request.limits.max_extracted_characters:
            raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.SCAN_LIMIT_EXCEEDED)
        self.extracted_characters = observed
        self.chunks.append(text)


class RegistryAttachmentDuplicateFingerprintFacade(AttachmentDuplicateFingerprintFacade):
    def __init__(
        self,
        *,
        resolver: AttachmentDuplicateFingerprintMaterialResolver,
    ) -> None:
        self._resolver = resolver
        self._cache: dict[
            str,
            tuple[str, AttachmentDuplicateFingerprintResultV2],
        ] = {}

    async def fingerprint(
        self,
        request: AttachmentDuplicateFingerprintRequestV2,
    ) -> AttachmentDuplicateFingerprintResultV2:
        validate_attachment_duplicate_fingerprint_request_identity(request)
        cached = self._cache.get(request.idempotency_key)
        if cached is not None:
            request_sha256, result = cached
            if request_sha256 != request.request_sha256:
                raise ValueError("attachment duplicate fingerprint idempotency conflict")
            return result

        try:
            path = self._resolver.resolve_output(request)
            validate_attachment_output_material(path, request)
            text = _extract_text(path, request)
            validate_attachment_output_material(path, request)
            tokens = _normalize_tokens(text)
            if len(tokens) < request.limits.min_token_count or len(tokens) < request.limits.shingle_size:
                raise _FingerprintUnsupported(
                    AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_INSUFFICIENT
                )
            fingerprint, shingle_count = _simhash(
                tokens,
                shingle_size=request.limits.shingle_size,
            )
            result = _result(
                request,
                outcome=AttachmentDuplicateFingerprintOutcomeV2.SUPPORTED,
                similarity_fingerprint=fingerprint,
                token_count=len(tokens),
                shingle_count=shingle_count,
                failure_code=None,
                extractor_version=ATTACHMENT_DUPLICATE_EXTRACTOR_VERSION,
            )
        except AttachmentDuplicateFingerprintMaterialError as exc:
            result = _result(
                request,
                outcome=AttachmentDuplicateFingerprintOutcomeV2.BLOCKED,
                failure_code=exc.code,
            )
        except _FingerprintUnsupported as exc:
            result = _result(
                request,
                outcome=AttachmentDuplicateFingerprintOutcomeV2.UNSUPPORTED,
                failure_code=exc.code,
                extractor_version=(
                    ATTACHMENT_DUPLICATE_EXTRACTOR_VERSION
                    if exc.code is AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_INSUFFICIENT
                    else None
                ),
            )
        except _FingerprintBlocked as exc:
            result = _result(
                request,
                outcome=AttachmentDuplicateFingerprintOutcomeV2.BLOCKED,
                failure_code=exc.code,
            )
        except Exception:
            result = _result(
                request,
                outcome=AttachmentDuplicateFingerprintOutcomeV2.BLOCKED,
                failure_code=(AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE),
            )
        self._cache[request.idempotency_key] = (
            request.request_sha256,
            result,
        )
        return result


def validate_attachment_output_material(
    path: Path,
    request: AttachmentTextExtractionRequest,
) -> None:
    if path.is_symlink():
        raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.SYMLINK_UNSUPPORTED)
    if not path.exists():
        raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.MATERIAL_NOT_FOUND)
    if not path.is_file() and not path.is_dir():
        raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE)
    if path.is_dir():
        size_bytes = 0
        file_identities: set[tuple[int, int]] = set()
        for child in path.rglob("*"):
            if child.is_symlink():
                raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.SYMLINK_UNSUPPORTED)
            if child.is_file():
                child_stat = child.stat()
                identity = (child_stat.st_dev, child_stat.st_ino)
                if identity in file_identities:
                    raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE)
                file_identities.add(identity)
                size_bytes += child_stat.st_size
    else:
        size_bytes = path.stat().st_size
    if size_bytes != request.size_bytes:
        raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.MATERIAL_MISMATCH)
    if sha256_path(path) != request.content_sha256:
        raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.OUTPUT_HASH_MISMATCH)


def _extract_text(
    path: Path,
    request: AttachmentTextExtractionRequest,
) -> str:
    chunks, _, _ = extract_attachment_text_projections(path, request)
    return "\n".join(chunks)


def extract_attachment_text_projections(
    path: Path,
    request: AttachmentTextExtractionRequest,
    *,
    require_all_members_scannable: bool = False,
) -> tuple[tuple[str, ...], tuple[str, ...], int]:
    state = _ExtractionState(
        request=request,
        require_all_members_scannable=require_all_members_scannable,
    )
    if path.is_dir():
        _extract_directory(path, state)
    else:
        _check_file_size(path.stat().st_size, request)
        suffix = path.suffix.casefold()
        if suffix in _UNSUPPORTED_CONTAINER_SUFFIXES:
            raise _FingerprintUnsupported(AttachmentDuplicateFingerprintFailureCodeV2.CONTAINER_UNSUPPORTED)
        if suffix in _TEXT_SUFFIXES or suffix in _STRUCTURED_DOCUMENT_SUFFIXES:
            state.add_member(size_bytes=path.stat().st_size)
            state.add_text(_extract_supported_path(path, request))
        elif suffix == ".zip" or zipfile.is_zipfile(path):
            if state.require_all_members_scannable:
                state.add_member(size_bytes=path.stat().st_size)
            _scan_zip_path(path, state, prefix="")
        else:
            raise _FingerprintUnsupported(AttachmentDuplicateFingerprintFailureCodeV2.MEDIA_UNSUPPORTED)
    return (
        tuple(state.chunks),
        tuple(state.paths),
        state.inventory_members,
    )


def _extract_directory(
    root: Path,
    state: _ExtractionState,
) -> None:
    for child in sorted(
        root.rglob("*"),
        key=lambda item: item.relative_to(root).as_posix(),
    ):
        logical_path = child.relative_to(root).as_posix()
        if child.is_symlink():
            raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.SYMLINK_UNSUPPORTED)
        if child.is_dir():
            state.add_member(
                size_bytes=0,
                logical_path=f"{logical_path}/",
            )
            continue
        if not child.is_file():
            raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE)
        size = child.stat().st_size
        _check_file_size(size, state.request)
        suffix = child.suffix.casefold()
        if suffix in _UNSUPPORTED_CONTAINER_SUFFIXES:
            raise _FingerprintUnsupported(AttachmentDuplicateFingerprintFailureCodeV2.CONTAINER_UNSUPPORTED)
        if suffix in _TEXT_SUFFIXES or suffix in _STRUCTURED_DOCUMENT_SUFFIXES:
            state.add_member(
                size_bytes=size,
                logical_path=logical_path,
            )
            state.add_text(_extract_supported_path(child, state.request))
        elif suffix == ".zip" or zipfile.is_zipfile(child):
            if state.require_all_members_scannable:
                state.add_member(
                    size_bytes=size,
                    logical_path=logical_path,
                )
            else:
                state.add_path(logical_path)
            _scan_zip_path(
                child,
                state,
                prefix=logical_path,
            )
        elif state.require_all_members_scannable:
            raise _FingerprintUnsupported(AttachmentDuplicateFingerprintFailureCodeV2.CONTAINER_UNSUPPORTED)


def _extract_supported_path(
    path: Path,
    request: AttachmentTextExtractionRequest,
) -> str:
    try:
        return extract_text(
            path,
            max_characters=request.limits.max_extracted_characters + 1,
        )
    except Exception as exc:
        raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE) from exc


def _scan_zip_path(
    path: Path,
    state: _ExtractionState,
    *,
    prefix: str,
) -> None:
    _check_file_size(path.stat().st_size, state.request)
    try:
        with path.open("rb") as handle:
            data = handle.read(state.request.limits.max_file_bytes + 1)
    except OSError as exc:
        raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE) from exc
    if len(data) > state.request.limits.max_file_bytes:
        raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.SCAN_LIMIT_EXCEEDED)
    _scan_zip_bytes(
        data,
        state=state,
        depth=1,
        prefix=prefix,
    )


def _scan_zip_bytes(
    data: bytes,
    *,
    state: _ExtractionState,
    depth: int,
    prefix: str,
) -> None:
    if depth > state.request.limits.max_nested_depth:
        raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.SCAN_LIMIT_EXCEEDED)
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise _FingerprintBlocked(
            AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE,
            reason=AttachmentTextExtractionFailureReason.CONTAINER_TRUNCATED,
        ) from exc
    seen: set[str] = set()
    try:
        for info in sorted(archive.infolist(), key=lambda item: item.filename):
            normalized = _normalize_archive_path(info.filename)
            logical_path = f"{prefix.rstrip('/')}/{normalized}" if prefix else normalized
            if normalized in seen:
                raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE)
            seen.add(normalized)
            if info.flag_bits & 0x1:
                raise _FingerprintBlocked(
                    AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE,
                    reason=AttachmentTextExtractionFailureReason.CONTAINER_ENCRYPTED,
                )
            mode = info.external_attr >> 16
            if stat.S_IFMT(mode) == stat.S_IFLNK:
                raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.SYMLINK_UNSUPPORTED)
            state.add_member(
                size_bytes=0 if info.is_dir() else info.file_size,
                logical_path=(f"{logical_path}/" if info.is_dir() else logical_path),
            )
            if info.is_dir():
                continue
            _check_file_size(info.file_size, state.request)
            ratio_milli = info.file_size * 1000 // max(info.compress_size, 1)
            if ratio_milli > state.request.limits.max_compression_ratio_milli:
                raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.SCAN_LIMIT_EXCEEDED)
            try:
                content = archive.read(info)
            except RuntimeError as exc:
                raise _FingerprintBlocked(
                    AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE,
                    reason=AttachmentTextExtractionFailureReason.CONTAINER_ENCRYPTED,
                ) from exc
            except (OSError, zipfile.BadZipFile) as exc:
                raise _FingerprintBlocked(
                    AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE,
                    reason=AttachmentTextExtractionFailureReason.CONTAINER_TRUNCATED,
                ) from exc
            suffix = PurePosixPath(normalized).suffix.casefold()
            if suffix in _TEXT_SUFFIXES or suffix in _STRUCTURED_DOCUMENT_SUFFIXES:
                text = _extract_supported_bytes(content, suffix=suffix)
                state.add_text(text)
            elif suffix == ".zip" or zipfile.is_zipfile(io.BytesIO(content)):
                _scan_zip_bytes(
                    content,
                    state=state,
                    depth=depth + 1,
                    prefix=logical_path,
                )
            elif state.require_all_members_scannable:
                raise _FingerprintUnsupported(
                    AttachmentDuplicateFingerprintFailureCodeV2.CONTAINER_UNSUPPORTED
                )
    finally:
        archive.close()


def _extract_supported_bytes(
    content: bytes,
    *,
    suffix: str,
) -> str:
    try:
        if suffix in _TEXT_SUFFIXES:
            return content.decode("utf-8", errors="replace")
        if suffix == ".docx":
            document = Document(io.BytesIO(content))
            docx_chunks = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                docx_chunks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            return "\n".join(docx_chunks)
        if suffix == ".xlsx":
            workbook = load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=False,
            )
            xlsx_chunks: list[str] = []
            try:
                for worksheet in workbook.worksheets:
                    xlsx_chunks.append(worksheet.title)
                    for row in worksheet.iter_rows(values_only=True):
                        xlsx_chunks.append(" | ".join("" if value is None else str(value) for value in row))
            finally:
                workbook.close()
            return "\n".join(xlsx_chunks)
        if suffix == ".pdf":
            with pymupdf.open(  # type: ignore[no-untyped-call]
                stream=content,
                filetype="pdf",
            ) as document:
                return "\n".join(page.get_text() for page in document)
    except Exception as exc:
        raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE) from exc
    return ""


def _normalize_archive_path(value: str) -> str:
    normalized = unicodedata.normalize(
        "NFKC",
        value.replace("\\", "/"),
    )
    path = PurePosixPath(normalized)
    if not normalized or normalized.startswith("/") or any(part in {"", ".", ".."} for part in path.parts):
        raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.OUTPUT_PATH_VIOLATION)
    return path.as_posix().rstrip("/")


def _check_file_size(
    size_bytes: int,
    request: AttachmentTextExtractionRequest,
) -> None:
    if size_bytes > request.limits.max_file_bytes:
        raise _FingerprintBlocked(AttachmentDuplicateFingerprintFailureCodeV2.SCAN_LIMIT_EXCEEDED)


def _normalize_tokens(text: str) -> tuple[str, ...]:
    normalized = unicodedata.normalize("NFKC", text).casefold()
    return tuple(match.group(0) for match in _TOKEN_PATTERN.finditer(normalized))


def _simhash(
    tokens: tuple[str, ...],
    *,
    shingle_size: int,
) -> tuple[str, int]:
    shingle_count = len(tokens) - shingle_size + 1
    weights = [0] * 256
    for index in range(shingle_count):
        shingle = "\x1f".join(tokens[index : index + shingle_size])
        digest = hashlib.sha256(shingle.encode("utf-8")).digest()
        for bit_index in range(256):
            byte = digest[bit_index // 8]
            mask = 1 << (7 - bit_index % 8)
            weights[bit_index] += 1 if byte & mask else -1
    value = 0
    for weight in weights:
        value = (value << 1) | int(weight >= 0)
    return value.to_bytes(32, byteorder="big").hex(), shingle_count


def _result(
    request: AttachmentDuplicateFingerprintRequestV2,
    *,
    outcome: AttachmentDuplicateFingerprintOutcomeV2,
    similarity_fingerprint: str | None = None,
    token_count: int = 0,
    shingle_count: int = 0,
    failure_code: AttachmentDuplicateFingerprintFailureCodeV2 | None,
    extractor_version: str | None = None,
) -> AttachmentDuplicateFingerprintResultV2:
    pending = AttachmentDuplicateFingerprintResultV2(
        fingerprint_result_id="attachment-duplicate-fingerprint-result://pending",
        fingerprint_request_ref=(attachment_duplicate_fingerprint_request_ref(request)),
        environment_artifact_ref=request.environment_artifact_ref,
        output_ref=request.output_ref,
        exact_content_sha256=request.content_sha256,
        outcome=outcome,
        similarity_fingerprint=similarity_fingerprint,
        token_count=token_count,
        shingle_count=shingle_count,
        failure_code=failure_code,
        extractor_version=extractor_version,
        normalization_version=ATTACHMENT_DUPLICATE_NORMALIZATION_VERSION,
        policy_version=ATTACHMENT_DUPLICATE_FINGERPRINT_POLICY_VERSION,
        result_sha256="0" * 64,
    )
    digest = attachment_duplicate_fingerprint_result_carried_sha256(pending)
    return pending.model_copy(
        update={
            "fingerprint_result_id": (f"attachment-duplicate-fingerprint-result://sha256/{digest}"),
            "result_sha256": digest,
        }
    )
