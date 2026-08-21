from __future__ import annotations

import asyncio
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook

from env_mock_agent.facade.contracts import FacadeObjectRef
from env_mock_agent.facade.cross_item_safety_adapter import (
    MappingAttachmentCrossItemSafetyMaterialResolver,
    RegistryAttachmentCrossItemSafetyScanFacade,
    StagingAttachmentCrossItemSafetyMaterialResolver,
)
from env_mock_agent.facade.cross_item_safety_v2 import (
    ATTACHMENT_CROSS_ITEM_SAFETY_POLICY_VERSION,
    AttachmentCrossItemSafetyScanFailureCodeV2,
    AttachmentCrossItemSafetyScanLimitsV2,
    AttachmentCrossItemSafetyScanRequestV2,
    AttachmentCrossItemSafetyScanStatusV2,
    attachment_cross_item_safety_scan_request_carried_sha256,
)
from env_mock_agent.facade.validation_v2 import (
    PROMPT_LEAKAGE_NORMALIZATION_VERSION,
    AttachmentValidationFingerprintCategoryV2,
    AttachmentValidationFingerprintMatchKindV2,
    AttachmentValidationFingerprintV2,
)
from env_mock_agent.providers.helpers import sha256_path

_TOKEN_PATTERN = re.compile(r"[\w]+(?:[._/-][\w]+)*", flags=re.UNICODE)


def _ref(
    object_type: str,
    suffix: str,
    *,
    digest: str = "a" * 64,
) -> FacadeObjectRef:
    return FacadeObjectRef(
        object_type=object_type,
        object_id=f"{object_type}://r7-02/{suffix}",
        object_version="v2",
        object_sha256=digest,
    )


def _tokens(value: str) -> tuple[str, ...]:
    normalized = unicodedata.normalize("NFKC", value).casefold()
    return tuple(match.group(0) for match in _TOKEN_PATTERN.finditer(normalized))


def _fingerprint(
    value: str,
    *,
    suffix: str,
    category: AttachmentValidationFingerprintCategoryV2 = (
        AttachmentValidationFingerprintCategoryV2.FINAL_ANSWER
    ),
    match_kind: AttachmentValidationFingerprintMatchKindV2 = (
        AttachmentValidationFingerprintMatchKindV2.TOKEN_WINDOW
    ),
) -> AttachmentValidationFingerprintV2:
    tokens = _tokens(value)
    digest = hashlib.sha256(" ".join(tokens).encode("utf-8")).hexdigest()
    return AttachmentValidationFingerprintV2(
        fingerprint_id=f"prompt-leakage-fingerprint://r7-02/{suffix}",
        category=category,
        match_kind=match_kind,
        digest_sha256=digest,
        token_count=len(tokens),
        normalization_version=PROMPT_LEAKAGE_NORMALIZATION_VERSION,
    )


def _limits(
    **updates: int,
) -> AttachmentCrossItemSafetyScanLimitsV2:
    values = {
        "max_file_bytes": 10_000_000,
        "max_extracted_characters": 1_000_000,
        "max_inventory_members": 10_000,
        "max_nested_depth": 4,
        "max_expanded_bytes": 50_000_000,
        "max_compression_ratio_milli": 1_000_000,
        "max_fingerprint_count": 10_000,
        "max_match_count": 1_000,
    }
    values.update(updates)
    return AttachmentCrossItemSafetyScanLimitsV2.model_validate(values)


def _request(
    path: Path,
    *,
    suffix: str,
    fingerprints: tuple[AttachmentValidationFingerprintV2, ...],
    logical_path: str | None = None,
    media_type: str = "text/plain",
    limits: AttachmentCrossItemSafetyScanLimitsV2 | None = None,
    content_sha256: str | None = None,
    idempotency_key: str | None = None,
) -> AttachmentCrossItemSafetyScanRequestV2:
    digest = content_sha256 or sha256_path(path)
    size_bytes = (
        path.stat().st_size
        if path.is_file()
        else sum(child.stat().st_size for child in path.rglob("*") if child.is_file())
    )
    ordered_fingerprints = tuple(
        sorted(
            fingerprints,
            key=lambda value: (
                value.category.value,
                value.match_kind.value,
                value.digest_sha256,
                value.token_count,
                value.fingerprint_id,
            ),
        )
    )
    pending = AttachmentCrossItemSafetyScanRequestV2(
        scan_request_id="attachment-cross-item-safety-scan-request://pending",
        target_item_id=f"dataset-item://r7-02/{suffix}",
        item_quality_result_ref=_ref("item-quality-compilation-result", suffix),
        environment_artifact_ref=_ref("environment-artifact", suffix),
        candidate_artifact_version_ref=_ref("candidate-artifact-version", suffix),
        output_ref=_ref("attachment-output", suffix, digest=digest),
        artifact_validation_result_ref=_ref(
            "artifact-deterministic-validation-result",
            suffix,
        ),
        logical_path=logical_path or f"inputs/{path.name}",
        media_type=media_type,
        content_sha256=digest,
        size_bytes=size_bytes,
        foreign_reference_set_refs=(_ref("prompt-leakage-reference-set", "foreign"),),
        foreign_fingerprints=ordered_fingerprints,
        limits=limits or _limits(),
        normalization_version=PROMPT_LEAKAGE_NORMALIZATION_VERSION,
        policy_version=ATTACHMENT_CROSS_ITEM_SAFETY_POLICY_VERSION,
        idempotency_key=(idempotency_key or f"attachment-cross-item-safety-idempotency://r7-02/{suffix}"),
        request_sha256="0" * 64,
    )
    request_digest = attachment_cross_item_safety_scan_request_carried_sha256(pending)
    return pending.model_copy(
        update={
            "scan_request_id": (f"attachment-cross-item-safety-scan-request://sha256/{request_digest}"),
            "request_sha256": request_digest,
        }
    )


def _facade(
    requests: tuple[
        tuple[AttachmentCrossItemSafetyScanRequestV2, Path],
        ...,
    ],
) -> RegistryAttachmentCrossItemSafetyScanFacade:
    return RegistryAttachmentCrossItemSafetyScanFacade(
        resolver=MappingAttachmentCrossItemSafetyMaterialResolver(
            outputs={request.output_ref.object_id: path for request, path in requests}
        )
    )


def _write_docx(path: Path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def _write_xlsx(path: Path, text: str) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Input"
    worksheet["A1"] = text
    workbook.save(path)
    workbook.close()


def _write_pdf(path: Path, text: str) -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(path)
    document.close()


def test_text_match_is_normalized_and_content_free(tmp_path: Path) -> None:
    path = tmp_path / "input.txt"
    path.write_text(
        ("\uff23\uff55\uff53\uff54\uff4f\uff4d\uff45\uff52 ALPHA secret final answer value is forty two"),
        encoding="utf-8",
    )
    fingerprint = _fingerprint(
        "customer alpha secret final answer value is forty two",
        suffix="answer",
    )
    request = _request(
        path,
        suffix="text",
        fingerprints=(fingerprint,),
    )

    result = asyncio.run(_facade(((request, path),)).scan(request))

    assert result.status is AttachmentCrossItemSafetyScanStatusV2.PASSED
    assert result.matched_fingerprint_ids == (fingerprint.fingerprint_id,)
    assert result.scan_complete is True
    serialized = json.dumps(result.model_dump(mode="json"), sort_keys=True)
    assert "customer alpha" not in serialized.casefold()
    assert str(tmp_path) not in serialized


@pytest.mark.parametrize(
    ("extension", "media_type", "writer"),
    (
        (
            ".json",
            "application/json",
            lambda path, text: path.write_text(
                json.dumps({"description": text}),
                encoding="utf-8",
            ),
        ),
        (
            ".csv",
            "text/csv",
            lambda path, text: path.write_text(
                f"name,description\nalpha,{text}\n",
                encoding="utf-8",
            ),
        ),
        (
            ".docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            _write_docx,
        ),
        (
            ".xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            _write_xlsx,
        ),
        (".pdf", "application/pdf", _write_pdf),
    ),
)
def test_supported_file_formats_match_foreign_fingerprint(
    tmp_path: Path,
    extension: str,
    media_type: str,
    writer,
) -> None:
    phrase = "foreign restricted answer contains unique quarterly total"
    path = tmp_path / f"input{extension}"
    writer(path, phrase)
    fingerprint = _fingerprint(phrase, suffix=extension.removeprefix("."))
    request = _request(
        path,
        suffix=extension.removeprefix("."),
        fingerprints=(fingerprint,),
        media_type=media_type,
    )

    result = asyncio.run(_facade(((request, path),)).scan(request))

    assert result.status is AttachmentCrossItemSafetyScanStatusV2.PASSED
    assert result.matched_fingerprint_ids == (fingerprint.fingerprint_id,)


def test_directory_and_nested_zip_scan_content_and_member_paths(
    tmp_path: Path,
) -> None:
    directory = tmp_path / "project"
    directory.mkdir()
    (directory / "input.txt").write_text(
        "ordinary unrelated input content",
        encoding="utf-8",
    )
    archive = directory / "bundle.zip"
    with zipfile.ZipFile(
        archive,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as outer:
        nested = io.BytesIO()
        with zipfile.ZipFile(
            nested,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as inner:
            inner.writestr(
                "answers/private-key.txt",
                "foreign restricted answer contains unique quarterly total",
            )
        outer.writestr("nested.zip", nested.getvalue())
    content_fingerprint = _fingerprint(
        "foreign restricted answer contains unique quarterly total",
        suffix="nested-content",
    )
    path_fingerprint = _fingerprint(
        "bundle.zip/nested.zip/answers/private-key.txt",
        suffix="nested-path",
        match_kind=AttachmentValidationFingerprintMatchKindV2.PATH_COMPONENT,
    )
    request = _request(
        directory,
        suffix="directory",
        fingerprints=(content_fingerprint, path_fingerprint),
        media_type="application/x-directory",
    )

    result = asyncio.run(_facade(((request, directory),)).scan(request))

    assert result.status is AttachmentCrossItemSafetyScanStatusV2.PASSED
    assert result.matched_fingerprint_ids == tuple(
        sorted(
            (
                content_fingerprint.fingerprint_id,
                path_fingerprint.fingerprint_id,
            )
        )
    )
    assert result.scanned_member_count >= 3


def test_clean_supported_attachment_returns_complete_empty_match(
    tmp_path: Path,
) -> None:
    path = tmp_path / "input.txt"
    path.write_text("clean unrelated input material", encoding="utf-8")
    request = _request(
        path,
        suffix="clean",
        fingerprints=(
            _fingerprint(
                "foreign restricted answer contains unique quarterly total",
                suffix="foreign",
            ),
        ),
    )

    result = asyncio.run(_facade(((request, path),)).scan(request))

    assert result.status is AttachmentCrossItemSafetyScanStatusV2.PASSED
    assert result.matched_fingerprint_ids == ()
    assert result.scan_complete is True


def test_unknown_directory_or_zip_member_blocks_complete_scan(
    tmp_path: Path,
) -> None:
    fingerprint = _fingerprint(
        "foreign restricted answer contains unique quarterly total",
        suffix="foreign",
    )
    directory = tmp_path / "project"
    directory.mkdir()
    (directory / "unknown.bin").write_bytes(b"\x00\x01")
    archive = tmp_path / "project.zip"
    with zipfile.ZipFile(
        archive,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as target:
        target.writestr("unknown.bin", b"\x00\x01")
    directory_request = _request(
        directory,
        suffix="unknown-directory",
        fingerprints=(fingerprint,),
        media_type="application/x-directory",
    )
    archive_request = _request(
        archive,
        suffix="unknown-archive",
        fingerprints=(fingerprint,),
        media_type="application/zip",
    )
    facade = _facade(
        (
            (directory_request, directory),
            (archive_request, archive),
        )
    )

    directory_result = asyncio.run(facade.scan(directory_request))
    archive_result = asyncio.run(facade.scan(archive_request))

    assert directory_result.failure_code is AttachmentCrossItemSafetyScanFailureCodeV2.CONTAINER_UNSUPPORTED
    assert archive_result.failure_code is AttachmentCrossItemSafetyScanFailureCodeV2.CONTAINER_UNSUPPORTED
    assert not directory_result.scan_complete
    assert not archive_result.scan_complete


def test_encrypted_and_truncated_zip_failures_are_precise(
    tmp_path: Path,
) -> None:
    encrypted = tmp_path / "encrypted.zip"
    with zipfile.ZipFile(encrypted, "w") as archive:
        archive.writestr("input.txt", "ordinary input material")
    encrypted_data = bytearray(encrypted.read_bytes())
    for signature in (b"PK\x03\x04", b"PK\x01\x02"):
        offset = 0
        while True:
            index = encrypted_data.find(signature, offset)
            if index < 0:
                break
            flag_offset = index + 6 if signature == b"PK\x03\x04" else index + 8
            flags = int.from_bytes(encrypted_data[flag_offset : flag_offset + 2], "little")
            encrypted_data[flag_offset : flag_offset + 2] = (flags | 1).to_bytes(2, "little")
            offset = index + 4
    encrypted.write_bytes(encrypted_data)

    truncated = tmp_path / "truncated.zip"
    with zipfile.ZipFile(truncated, "w") as archive:
        archive.writestr("input.txt", "ordinary input material")
    truncated.write_bytes(truncated.read_bytes()[:-8])

    fingerprint = _fingerprint(
        "foreign restricted answer contains unique quarterly total",
        suffix="foreign-container",
    )
    encrypted_request = _request(
        encrypted,
        suffix="encrypted",
        fingerprints=(fingerprint,),
        media_type="application/zip",
    )
    truncated_request = _request(
        truncated,
        suffix="truncated",
        fingerprints=(fingerprint,),
        media_type="application/zip",
    )
    facade = _facade(
        (
            (encrypted_request, encrypted),
            (truncated_request, truncated),
        )
    )

    encrypted_result = asyncio.run(facade.scan(encrypted_request))
    truncated_result = asyncio.run(facade.scan(truncated_request))

    assert encrypted_result.failure_code is (AttachmentCrossItemSafetyScanFailureCodeV2.CONTAINER_ENCRYPTED)
    assert truncated_result.failure_code is (AttachmentCrossItemSafetyScanFailureCodeV2.CONTAINER_TRUNCATED)
    assert not encrypted_result.matched_fingerprint_ids
    assert not truncated_result.matched_fingerprint_ids


def test_unsupported_hash_symlink_and_limits_fail_closed(
    tmp_path: Path,
) -> None:
    phrase = "foreign restricted answer contains unique quarterly total"
    fingerprint = _fingerprint(phrase, suffix="answer")
    binary = tmp_path / "input.bin"
    binary.write_bytes(b"\x00\x01\x02\x03")
    text = tmp_path / "input.txt"
    text.write_text(phrase, encoding="utf-8")
    link = tmp_path / "link.txt"
    link.symlink_to(text)
    unsupported = _request(
        binary,
        suffix="unsupported",
        fingerprints=(fingerprint,),
        media_type="application/octet-stream",
    )
    mismatch = _request(
        text,
        suffix="mismatch",
        fingerprints=(fingerprint,),
        content_sha256="b" * 64,
    )
    linked = _request(
        link,
        suffix="linked",
        fingerprints=(fingerprint,),
    )
    constrained = _request(
        text,
        suffix="constrained",
        fingerprints=(fingerprint,),
        limits=_limits(max_file_bytes=4),
    )
    facade = _facade(
        (
            (unsupported, binary),
            (mismatch, text),
            (linked, link),
            (constrained, text),
        )
    )

    results = tuple(
        asyncio.run(facade.scan(request)) for request in (unsupported, mismatch, linked, constrained)
    )

    assert tuple(result.status for result in results) == (AttachmentCrossItemSafetyScanStatusV2.BLOCKED,) * 4
    assert tuple(result.failure_code for result in results) == (
        AttachmentCrossItemSafetyScanFailureCodeV2.CONTAINER_UNSUPPORTED,
        AttachmentCrossItemSafetyScanFailureCodeV2.OUTPUT_HASH_MISMATCH,
        AttachmentCrossItemSafetyScanFailureCodeV2.SYMLINK_UNSUPPORTED,
        AttachmentCrossItemSafetyScanFailureCodeV2.SCAN_LIMIT_EXCEEDED,
    )
    assert all(not result.matched_fingerprint_ids for result in results)


def test_match_budget_and_idempotency_conflict_are_closed(
    tmp_path: Path,
) -> None:
    phrase = "foreign restricted answer contains unique quarterly total"
    path = tmp_path / "input.txt"
    path.write_text(phrase, encoding="utf-8")
    fingerprints = (
        _fingerprint(phrase, suffix="answer-a"),
        _fingerprint(
            phrase,
            suffix="answer-b",
            category=AttachmentValidationFingerprintCategoryV2.PRIVATE_REFERENCE,
        ),
    )
    constrained = _request(
        path,
        suffix="match-budget",
        fingerprints=fingerprints,
        limits=_limits(max_match_count=1),
    )
    first = _request(
        path,
        suffix="first",
        fingerprints=(fingerprints[0],),
        idempotency_key="attachment-cross-item-safety-idempotency://shared",
    )
    changed = _request(
        path,
        suffix="changed",
        fingerprints=(fingerprints[1],),
        idempotency_key="attachment-cross-item-safety-idempotency://shared",
    )
    facade = _facade(
        (
            (constrained, path),
            (first, path),
            (changed, path),
        )
    )

    constrained_result = asyncio.run(facade.scan(constrained))
    first_result = asyncio.run(facade.scan(first))
    conflict_result = asyncio.run(facade.scan(changed))

    assert constrained_result.status is AttachmentCrossItemSafetyScanStatusV2.BLOCKED
    assert constrained_result.failure_code is (AttachmentCrossItemSafetyScanFailureCodeV2.SCAN_LIMIT_EXCEEDED)
    assert first_result.status is AttachmentCrossItemSafetyScanStatusV2.PASSED
    assert conflict_result.status is AttachmentCrossItemSafetyScanStatusV2.BLOCKED
    assert conflict_result.failure_code is (AttachmentCrossItemSafetyScanFailureCodeV2.IDEMPOTENCY_CONFLICT)


def test_strict_zip_container_counts_against_inventory_budget(
    tmp_path: Path,
) -> None:
    phrase = "foreign restricted answer contains unique quarterly total"
    archive = tmp_path / "input.zip"
    with zipfile.ZipFile(
        archive,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as target:
        target.writestr("input.txt", phrase)
    fingerprint = _fingerprint(phrase, suffix="inventory")
    request = _request(
        archive,
        suffix="inventory",
        fingerprints=(fingerprint,),
        media_type="application/zip",
        limits=_limits(max_inventory_members=1),
    )

    result = asyncio.run(_facade(((request, archive),)).scan(request))

    assert result.status is AttachmentCrossItemSafetyScanStatusV2.BLOCKED
    assert result.failure_code is (AttachmentCrossItemSafetyScanFailureCodeV2.SCAN_LIMIT_EXCEEDED)


def test_staging_resolver_rejects_path_escape(tmp_path: Path) -> None:
    staging = tmp_path / "staging"
    staging.mkdir()
    outside = tmp_path / "outside.txt"
    outside.write_text(
        "foreign restricted answer contains unique quarterly total",
        encoding="utf-8",
    )
    fingerprint = _fingerprint(outside.read_text(), suffix="outside")
    request = _request(
        outside,
        suffix="escape",
        fingerprints=(fingerprint,),
    )
    facade = RegistryAttachmentCrossItemSafetyScanFacade(
        resolver=StagingAttachmentCrossItemSafetyMaterialResolver(
            staging_root=staging,
            outputs={request.output_ref.object_id: outside},
        )
    )

    result = asyncio.run(facade.scan(request))

    assert result.status is AttachmentCrossItemSafetyScanStatusV2.BLOCKED
    assert result.failure_code is (AttachmentCrossItemSafetyScanFailureCodeV2.OUTPUT_PATH_VIOLATION)
