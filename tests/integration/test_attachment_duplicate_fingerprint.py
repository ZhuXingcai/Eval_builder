from __future__ import annotations

import asyncio
import io
import json
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook

from env_mock_agent.facade.contracts import FacadeObjectRef
from env_mock_agent.facade.duplicate_adapter import (
    MappingAttachmentDuplicateFingerprintMaterialResolver,
    RegistryAttachmentDuplicateFingerprintFacade,
    StagingAttachmentDuplicateFingerprintMaterialResolver,
)
from env_mock_agent.facade.duplicate_v2 import (
    ATTACHMENT_DUPLICATE_FINGERPRINT_POLICY_VERSION,
    ATTACHMENT_DUPLICATE_NORMALIZATION_VERSION,
    AttachmentDuplicateFingerprintFailureCodeV2,
    AttachmentDuplicateFingerprintLimitsV2,
    AttachmentDuplicateFingerprintOutcomeV2,
    AttachmentDuplicateFingerprintRequestV2,
    attachment_duplicate_fingerprint_request_carried_sha256,
)
from env_mock_agent.providers.helpers import sha256_path


def _ref(
    object_type: str,
    suffix: str,
    *,
    digest: str = "a" * 64,
) -> FacadeObjectRef:
    return FacadeObjectRef(
        object_type=object_type,
        object_id=f"{object_type}://r7-01/{suffix}",
        object_version="v2",
        object_sha256=digest,
    )


def _limits(
    **updates: int,
) -> AttachmentDuplicateFingerprintLimitsV2:
    values = {
        "max_file_bytes": 10_000_000,
        "max_extracted_characters": 1_000_000,
        "max_inventory_members": 10_000,
        "max_nested_depth": 4,
        "max_expanded_bytes": 50_000_000,
        "max_compression_ratio_milli": 1_000_000,
        "min_token_count": 4,
        "shingle_size": 3,
        "fingerprint_bits": 256,
    }
    values.update(updates)
    return AttachmentDuplicateFingerprintLimitsV2.model_validate(values)


def _request(
    path: Path,
    *,
    suffix: str,
    logical_path: str | None = None,
    media_type: str = "text/plain",
    limits: AttachmentDuplicateFingerprintLimitsV2 | None = None,
    content_sha256: str | None = None,
    idempotency_key: str | None = None,
) -> AttachmentDuplicateFingerprintRequestV2:
    digest = content_sha256 or sha256_path(path)
    size_bytes = (
        path.stat().st_size
        if path.is_file()
        else sum(child.stat().st_size for child in path.rglob("*") if child.is_file())
    )
    pending = AttachmentDuplicateFingerprintRequestV2(
        fingerprint_request_id="attachment-duplicate-fingerprint-request://pending",
        item_quality_result_ref=_ref(
            "item-quality-compilation-result",
            suffix,
        ),
        environment_artifact_ref=_ref("environment-artifact", suffix),
        candidate_artifact_version_ref=_ref(
            "candidate-artifact-version",
            suffix,
        ),
        output_ref=_ref("attachment-output", suffix, digest=digest),
        artifact_validation_result_ref=_ref(
            "artifact-deterministic-validation-result",
            suffix,
        ),
        logical_path=logical_path or f"inputs/{path.name}",
        media_type=media_type,
        content_sha256=digest,
        size_bytes=size_bytes,
        limits=limits or _limits(),
        normalization_version=ATTACHMENT_DUPLICATE_NORMALIZATION_VERSION,
        policy_version=ATTACHMENT_DUPLICATE_FINGERPRINT_POLICY_VERSION,
        idempotency_key=(idempotency_key or f"attachment-duplicate-fingerprint-idempotency://{suffix}"),
        request_sha256="0" * 64,
    )
    request_digest = attachment_duplicate_fingerprint_request_carried_sha256(pending)
    return pending.model_copy(
        update={
            "fingerprint_request_id": (f"attachment-duplicate-fingerprint-request://sha256/{request_digest}"),
            "request_sha256": request_digest,
        }
    )


def _facade(
    requests: tuple[
        tuple[AttachmentDuplicateFingerprintRequestV2, Path],
        ...,
    ],
) -> RegistryAttachmentDuplicateFingerprintFacade:
    return RegistryAttachmentDuplicateFingerprintFacade(
        resolver=MappingAttachmentDuplicateFingerprintMaterialResolver(
            outputs={request.output_ref.object_id: path for request, path in requests}
        )
    )


def _write_docx(path: Path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def _write_xlsx(path: Path, text: str) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Input"
    worksheet["A1"] = text
    workbook.save(path)
    workbook.close()


def _write_pdf(path: Path, text: str) -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(path)
    document.close()


def test_text_normalization_is_stable_and_content_free(
    tmp_path: Path,
) -> None:
    left = tmp_path / "left.txt"
    right = tmp_path / "right.txt"
    left.write_text(
        ("\uff23\uff55\uff53\uff54\uff4f\uff4d\uff45\uff52   ALPHA needs a deterministic report.\n"),
        encoding="utf-8",
    )
    right.write_text(
        "customer alpha NEEDS a deterministic report.",
        encoding="utf-8",
    )
    left_request = _request(left, suffix="left")
    right_request = _request(right, suffix="right")
    facade = _facade(
        (
            (left_request, left),
            (right_request, right),
        )
    )

    left_result, right_result = asyncio.run(
        _fingerprint_pair(
            facade,
            left_request,
            right_request,
        )
    )

    assert left_result.outcome is AttachmentDuplicateFingerprintOutcomeV2.SUPPORTED
    assert right_result.outcome is AttachmentDuplicateFingerprintOutcomeV2.SUPPORTED
    assert left_result.similarity_fingerprint == (right_result.similarity_fingerprint)
    assert left_result.token_count == right_result.token_count
    serialized = json.dumps(left_result.model_dump(mode="json"), sort_keys=True)
    assert "customer" not in serialized.casefold()
    assert str(tmp_path) not in serialized


async def _fingerprint_pair(
    facade: RegistryAttachmentDuplicateFingerprintFacade,
    left: AttachmentDuplicateFingerprintRequestV2,
    right: AttachmentDuplicateFingerprintRequestV2,
):
    return await facade.fingerprint(left), await facade.fingerprint(right)


@pytest.mark.parametrize(
    ("extension", "media_type", "writer"),
    (
        (
            ".json",
            "application/json",
            lambda path, text: path.write_text(
                json.dumps({"description": text}),
                encoding="utf-8",
            ),
        ),
        (
            ".csv",
            "text/csv",
            lambda path, text: path.write_text(
                f"name,description\nalpha,{text}\n",
                encoding="utf-8",
            ),
        ),
        (".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", _write_docx),
        (".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _write_xlsx),
        (".pdf", "application/pdf", _write_pdf),
    ),
)
def test_supported_file_formats_are_fingerprinted(
    tmp_path: Path,
    extension: str,
    media_type: str,
    writer,
) -> None:
    path = tmp_path / f"input{extension}"
    writer(
        path,
        "customer alpha requires a deterministic quarterly report",
    )
    request = _request(
        path,
        suffix=extension.removeprefix("."),
        media_type=media_type,
    )
    result = asyncio.run(_facade(((request, path),)).fingerprint(request))

    assert result.outcome is AttachmentDuplicateFingerprintOutcomeV2.SUPPORTED
    assert result.similarity_fingerprint is not None
    assert result.token_count >= request.limits.min_token_count


def test_directory_and_generic_zip_aggregate_text_members(
    tmp_path: Path,
) -> None:
    directory = tmp_path / "project"
    directory.mkdir()
    (directory / "a.txt").write_text(
        "customer alpha quarterly report requirements",
        encoding="utf-8",
    )
    (directory / "b.json").write_text(
        '{"status":"draft input material"}',
        encoding="utf-8",
    )
    archive = tmp_path / "project.zip"
    with zipfile.ZipFile(
        archive,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as target:
        target.writestr(
            "a.txt",
            "customer alpha quarterly report requirements",
        )
        target.writestr(
            "b.json",
            '{"status":"draft input material"}',
        )
    directory_request = _request(
        directory,
        suffix="directory",
        media_type="application/x-directory",
    )
    archive_request = _request(
        archive,
        suffix="archive",
        media_type="application/zip",
    )
    facade = _facade(
        (
            (directory_request, directory),
            (archive_request, archive),
        )
    )

    directory_result, archive_result = asyncio.run(
        _fingerprint_pair(
            facade,
            directory_request,
            archive_request,
        )
    )

    assert directory_result.outcome is AttachmentDuplicateFingerprintOutcomeV2.SUPPORTED
    assert archive_result.outcome is AttachmentDuplicateFingerprintOutcomeV2.SUPPORTED
    assert directory_result.similarity_fingerprint == (archive_result.similarity_fingerprint)


def test_unsupported_and_insufficient_content_are_explicit(
    tmp_path: Path,
) -> None:
    binary = tmp_path / "input.bin"
    binary.write_bytes(b"\x00\x01\x02\x03")
    tiny = tmp_path / "tiny.txt"
    tiny.write_text("too short", encoding="utf-8")
    binary_request = _request(
        binary,
        suffix="binary",
        media_type="application/octet-stream",
    )
    tiny_request = _request(
        tiny,
        suffix="tiny",
        limits=_limits(min_token_count=5),
    )
    facade = _facade(
        (
            (binary_request, binary),
            (tiny_request, tiny),
        )
    )

    binary_result, tiny_result = asyncio.run(
        _fingerprint_pair(
            facade,
            binary_request,
            tiny_request,
        )
    )

    assert binary_result.outcome is AttachmentDuplicateFingerprintOutcomeV2.UNSUPPORTED
    assert binary_result.failure_code is (AttachmentDuplicateFingerprintFailureCodeV2.MEDIA_UNSUPPORTED)
    assert tiny_result.outcome is AttachmentDuplicateFingerprintOutcomeV2.UNSUPPORTED
    assert tiny_result.failure_code is (AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_INSUFFICIENT)


def test_hash_symlink_and_scan_limits_fail_closed(
    tmp_path: Path,
) -> None:
    path = tmp_path / "input.txt"
    path.write_text(
        "customer alpha quarterly report requirements",
        encoding="utf-8",
    )
    mismatch = _request(
        path,
        suffix="mismatch",
        content_sha256="b" * 64,
    )
    constrained = _request(
        path,
        suffix="constrained",
        limits=_limits(max_file_bytes=4),
    )
    link = tmp_path / "link.txt"
    link.symlink_to(path)
    linked = _request(link, suffix="linked")
    facade = _facade(
        (
            (mismatch, path),
            (constrained, path),
            (linked, link),
        )
    )

    mismatch_result = asyncio.run(facade.fingerprint(mismatch))
    constrained_result = asyncio.run(facade.fingerprint(constrained))
    linked_result = asyncio.run(facade.fingerprint(linked))

    assert mismatch_result.failure_code is (AttachmentDuplicateFingerprintFailureCodeV2.OUTPUT_HASH_MISMATCH)
    assert constrained_result.failure_code is (
        AttachmentDuplicateFingerprintFailureCodeV2.SCAN_LIMIT_EXCEEDED
    )
    assert linked_result.failure_code is (AttachmentDuplicateFingerprintFailureCodeV2.SYMLINK_UNSUPPORTED)
    assert all(
        result.outcome is AttachmentDuplicateFingerprintOutcomeV2.BLOCKED
        for result in (
            mismatch_result,
            constrained_result,
            linked_result,
        )
    )


def test_zip_member_and_compression_limits_fail_closed(
    tmp_path: Path,
) -> None:
    archive = tmp_path / "many.zip"
    with zipfile.ZipFile(
        archive,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as target:
        target.writestr("a.txt", "alpha report requirements")
        target.writestr("b.txt", "beta report requirements")
    request = _request(
        archive,
        suffix="many",
        media_type="application/zip",
        limits=_limits(max_inventory_members=1),
    )
    result = asyncio.run(_facade(((request, archive),)).fingerprint(request))

    assert result.outcome is AttachmentDuplicateFingerprintOutcomeV2.BLOCKED
    assert result.failure_code is (AttachmentDuplicateFingerprintFailureCodeV2.SCAN_LIMIT_EXCEEDED)


def test_encrypted_and_truncated_zip_remain_content_unscannable(
    tmp_path: Path,
) -> None:
    encrypted = tmp_path / "encrypted.zip"
    with zipfile.ZipFile(encrypted, "w") as archive:
        archive.writestr("input.txt", "ordinary input material")
    encrypted_data = bytearray(encrypted.read_bytes())
    for signature in (b"PK\x03\x04", b"PK\x01\x02"):
        offset = 0
        while True:
            index = encrypted_data.find(signature, offset)
            if index < 0:
                break
            flag_offset = index + 6 if signature == b"PK\x03\x04" else index + 8
            flags = int.from_bytes(encrypted_data[flag_offset : flag_offset + 2], "little")
            encrypted_data[flag_offset : flag_offset + 2] = (flags | 1).to_bytes(2, "little")
            offset = index + 4
    encrypted.write_bytes(encrypted_data)

    truncated = tmp_path / "truncated.zip"
    with zipfile.ZipFile(truncated, "w") as archive:
        archive.writestr("input.txt", "ordinary input material")
    truncated.write_bytes(truncated.read_bytes()[:-8])

    encrypted_request = _request(
        encrypted,
        suffix="encrypted",
        media_type="application/zip",
    )
    truncated_request = _request(
        truncated,
        suffix="truncated",
        media_type="application/zip",
    )
    facade = _facade(
        (
            (encrypted_request, encrypted),
            (truncated_request, truncated),
        )
    )

    encrypted_result = asyncio.run(facade.fingerprint(encrypted_request))
    truncated_result = asyncio.run(facade.fingerprint(truncated_request))

    assert encrypted_result.failure_code is (AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE)
    assert truncated_result.failure_code is (AttachmentDuplicateFingerprintFailureCodeV2.CONTENT_UNSCANNABLE)


def test_idempotency_replay_and_conflict(
    tmp_path: Path,
) -> None:
    left = tmp_path / "left.txt"
    right = tmp_path / "right.txt"
    left.write_text(
        "customer alpha quarterly report requirements",
        encoding="utf-8",
    )
    right.write_text(
        "customer beta annual report requirements",
        encoding="utf-8",
    )
    key = "attachment-duplicate-fingerprint-idempotency://shared"
    left_request = _request(
        left,
        suffix="left",
        idempotency_key=key,
    )
    right_request = _request(
        right,
        suffix="right",
        idempotency_key=key,
    )
    facade = _facade(
        (
            (left_request, left),
            (right_request, right),
        )
    )

    first = asyncio.run(facade.fingerprint(left_request))
    replay = asyncio.run(facade.fingerprint(left_request))
    assert replay is first
    with pytest.raises(ValueError, match="idempotency conflict"):
        asyncio.run(facade.fingerprint(right_request))


def test_missing_material_and_staging_escape_are_blocked(
    tmp_path: Path,
) -> None:
    outside = tmp_path / "outside.txt"
    outside.write_text(
        "customer alpha quarterly report requirements",
        encoding="utf-8",
    )
    missing = _request(outside, suffix="missing")
    escaped = _request(outside, suffix="escaped")
    missing_facade = RegistryAttachmentDuplicateFingerprintFacade(
        resolver=MappingAttachmentDuplicateFingerprintMaterialResolver(outputs={})
    )
    staging_root = tmp_path / "staging"
    staging_root.mkdir()
    escaped_facade = RegistryAttachmentDuplicateFingerprintFacade(
        resolver=StagingAttachmentDuplicateFingerprintMaterialResolver(
            staging_root=staging_root,
            outputs={escaped.output_ref.object_id: outside},
        )
    )

    missing_result = asyncio.run(missing_facade.fingerprint(missing))
    escaped_result = asyncio.run(escaped_facade.fingerprint(escaped))

    assert missing_result.failure_code is (AttachmentDuplicateFingerprintFailureCodeV2.MATERIAL_NOT_FOUND)
    assert escaped_result.failure_code is (AttachmentDuplicateFingerprintFailureCodeV2.OUTPUT_PATH_VIOLATION)


def test_character_expanded_byte_and_container_limits(
    tmp_path: Path,
) -> None:
    text = tmp_path / "large.txt"
    text.write_text(
        "customer alpha quarterly report requirements",
        encoding="utf-8",
    )
    character_request = _request(
        text,
        suffix="character-limit",
        limits=_limits(max_extracted_characters=8),
    )
    expanded_request = _request(
        text,
        suffix="expanded-limit",
        limits=_limits(max_expanded_bytes=4),
    )
    unsupported = tmp_path / "archive.tar"
    unsupported.write_text(
        "customer alpha quarterly report requirements",
        encoding="utf-8",
    )
    unsupported_request = _request(
        unsupported,
        suffix="unsupported-container",
        media_type="application/x-tar",
    )
    facade = _facade(
        (
            (character_request, text),
            (expanded_request, text),
            (unsupported_request, unsupported),
        )
    )

    character_result = asyncio.run(facade.fingerprint(character_request))
    expanded_result = asyncio.run(facade.fingerprint(expanded_request))
    unsupported_result = asyncio.run(facade.fingerprint(unsupported_request))

    assert character_result.failure_code is (AttachmentDuplicateFingerprintFailureCodeV2.SCAN_LIMIT_EXCEEDED)
    assert expanded_result.failure_code is (AttachmentDuplicateFingerprintFailureCodeV2.SCAN_LIMIT_EXCEEDED)
    assert unsupported_result.outcome is (AttachmentDuplicateFingerprintOutcomeV2.UNSUPPORTED)
    assert unsupported_result.failure_code is (
        AttachmentDuplicateFingerprintFailureCodeV2.CONTAINER_UNSUPPORTED
    )


def test_zip_traversal_depth_and_compression_limits(
    tmp_path: Path,
) -> None:
    traversal = tmp_path / "traversal.zip"
    with zipfile.ZipFile(traversal, "w") as target:
        target.writestr("../escape.txt", "unsafe archive path")

    inner_bytes = io.BytesIO()
    with zipfile.ZipFile(inner_bytes, "w") as inner:
        inner.writestr(
            "input.txt",
            "customer alpha quarterly report requirements",
        )
    nested = tmp_path / "nested.zip"
    with zipfile.ZipFile(nested, "w") as outer:
        outer.writestr("inner.zip", inner_bytes.getvalue())

    compressed = tmp_path / "compressed.zip"
    with zipfile.ZipFile(
        compressed,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as target:
        target.writestr("large.txt", "a" * 10_000)

    traversal_request = _request(
        traversal,
        suffix="traversal",
        media_type="application/zip",
    )
    nested_request = _request(
        nested,
        suffix="nested-depth",
        media_type="application/zip",
        limits=_limits(max_nested_depth=1),
    )
    compressed_request = _request(
        compressed,
        suffix="compression",
        media_type="application/zip",
        limits=_limits(max_compression_ratio_milli=1_000),
    )
    facade = _facade(
        (
            (traversal_request, traversal),
            (nested_request, nested),
            (compressed_request, compressed),
        )
    )

    traversal_result = asyncio.run(facade.fingerprint(traversal_request))
    nested_result = asyncio.run(facade.fingerprint(nested_request))
    compressed_result = asyncio.run(facade.fingerprint(compressed_request))

    assert traversal_result.failure_code is (
        AttachmentDuplicateFingerprintFailureCodeV2.OUTPUT_PATH_VIOLATION
    )
    assert nested_result.failure_code is (AttachmentDuplicateFingerprintFailureCodeV2.SCAN_LIMIT_EXCEEDED)
    assert compressed_result.failure_code is (AttachmentDuplicateFingerprintFailureCodeV2.SCAN_LIMIT_EXCEEDED)


def test_zip_structured_members_are_bounded_and_supported(
    tmp_path: Path,
) -> None:
    docx = tmp_path / "nested.docx"
    xlsx = tmp_path / "nested.xlsx"
    pdf = tmp_path / "nested.pdf"
    text = "customer alpha quarterly report requirements"
    _write_docx(docx, text)
    _write_xlsx(xlsx, text)
    _write_pdf(pdf, text)
    archive = tmp_path / "structured.zip"
    with zipfile.ZipFile(
        archive,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as target:
        target.writestr("documents/", b"")
        target.writestr("documents/input.docx", docx.read_bytes())
        target.writestr("documents/input.xlsx", xlsx.read_bytes())
        target.writestr("documents/input.pdf", pdf.read_bytes())
    request = _request(
        archive,
        suffix="structured-members",
        media_type="application/zip",
    )

    result = asyncio.run(_facade(((request, archive),)).fingerprint(request))

    assert result.outcome is (AttachmentDuplicateFingerprintOutcomeV2.SUPPORTED)
    assert result.similarity_fingerprint is not None
